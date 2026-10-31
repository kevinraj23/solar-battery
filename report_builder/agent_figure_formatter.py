"""
agent_figure_formatter.py
=========================
Manages high-resolution component image insertion and clean, formatted placeholder
boxes for non-component diagrams (Block Diagram, Circuit Diagram, Hardware Kit,
Flowchart, and Working Model Photo) matching the sample PDF.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_NAME = "Times New Roman"
CAPTION_FONT_SIZE = Pt(11)

def add_component_image(doc, image_path, caption_text, width_inches=3.0):
    """
    Inserts a high-quality component product image centered on the page
    followed by a standard figure caption.
    """
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
    else:
        # Fallback if image not found on disk
        add_image_placeholder(doc, caption_text, caption_text, height_inches=2.2)
        return

    # Add Caption
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = FONT_NAME
    r_cap.font.size = CAPTION_FONT_SIZE
    r_cap.font.bold = True

def add_image_placeholder(doc, box_title, caption_text, height_inches=2.8):
    """
    Creates a clean, styled, light-grey dashed bordered box spanning the width of the page
    where the user will insert their own diagram or photo, followed by the caption.
    """
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(5.8)
    tbl.rows[0].height = Inches(height_inches)
    
    # Set light dashed border via Word XML
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="dashed" w:sz="6" w:space="0" w:color="CBD5E0"/>
            <w:left w:val="dashed" w:sz="6" w:space="0" w:color="CBD5E0"/>
            <w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="CBD5E0"/>
            <w:right w:val="dashed" w:sz="6" w:space="0" w:color="CBD5E0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    # Set very light background shading
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    tcPr.append(shd)
    
    # Prompt text inside cell
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(int(height_inches * 24))
    p.paragraph_format.space_after = Pt(4)
    
    r1 = p.add_run(f"📷  [ INSERT {box_title.upper()} HERE ]\n")
    r1.font.name = FONT_NAME
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(74, 85, 104)
    
    r2 = p.add_run("(Leave space as requested: Replace this box with your actual schematic / diagram / photo)")
    r2.font.name = FONT_NAME
    r2.font.size = Pt(9.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(160, 174, 192)
    
    # Figure Caption below placeholder box
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(6)
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = FONT_NAME
    r_cap.font.size = CAPTION_FONT_SIZE
    r_cap.font.bold = True
