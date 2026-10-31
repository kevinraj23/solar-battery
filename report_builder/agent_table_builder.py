"""
agent_table_builder.py
======================
Builds precisely styled tables for Table of Contents, List of Figures, Test Results,
and Cost Estimation matching the visual format of the sample PDF.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_NAME = "Times New Roman"

def style_table_borders(table):
    """Applies clean, solid black borders to all table cells matching academic report formatting."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_cell_p(p, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.alignment = align
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)

def build_toc_table_page1(doc):
    """Generates the first page of the Table of Contents matching page vi of reference PDF."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_borders(table)

    headers = ["CHAP.NO.", "TITLE", "PAGE.NO"]
    col_widths = [Inches(1.2), Inches(3.8), Inches(1.0)]

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(title)
        r.font.name = FONT_NAME
        r.font.size = Pt(10.5)
        r.font.bold = True

    entries = [
        ("", "ABSTRACT", "v"),
        ("", "LIST OF FIGURES", "viii"),
        ("", "LIST OF ABBREVIATIONS", "1"),
        ("1", "INTRODUCTION", "2"),
        ("", "1.1   GENERAL", "2"),
        ("", "1.2   OBJECTIVES", "2"),
        ("", "1.3   PROBLEM STATEMENT", "3"),
        ("", "1.4   DISADVANTAGES OF EXISTING SYSTEM", "4"),
        ("", "1.5   PROPOSED SYSTEM", "4"),
        ("2", "LITERATURE SURVEY", "6"),
        ("", "2.1   HOME AUTOMATION-BASED ENVIRONMENTAL ALERT SYSTEM", "6"),
        ("", "2.2   DYNAMIC LIGHTING AND AR SMART HOME SYSTEM", "6"),
        ("", "2.3   SMART EMERGENCY SAFETY ALERT SYSTEM", "7"),
        ("", "2.4   IOT-BASED GAS AND ENVIRONMENTAL MONITORING", "7"),
        ("", "2.5   INCLUSIVE EMERGENCY ALERTS USING CLOUD NOTIFICATIONS", "8"),
        ("3", "SYSTEM CONSIDERATION", "9"),
        ("", "3.1   CIRCUIT DIAGRAM", "12"),
        ("", "3.2   HARDWARE SPECIFICATIONS", "14"),
    ]

    for chap, title, page in entries:
        row = table.add_row()
        cells = row.cells
        for i in range(3):
            cells[i].width = col_widths[i]
            
        p0 = cells[0].paragraphs[0]
        format_cell_p(p0, WD_ALIGN_PARAGRAPH.CENTER)
        r0 = p0.add_run(chap)
        r0.font.name = FONT_NAME
        r0.font.size = Pt(9.5)
        r0.font.bold = bool(chap)

        p1 = cells[1].paragraphs[0]
        format_cell_p(p1, WD_ALIGN_PARAGRAPH.LEFT)
        r1 = p1.add_run(title)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(9.5)
        r1.font.bold = bool(chap) or title in ["ABSTRACT", "LIST OF FIGURES", "LIST OF ABBREVIATIONS"]

        p2 = cells[2].paragraphs[0]
        format_cell_p(p2, WD_ALIGN_PARAGRAPH.CENTER)
        r2 = p2.add_run(page)
        r2.font.name = FONT_NAME
        r2.font.size = Pt(9.5)
        r2.font.bold = bool(chap)

def build_toc_table_page2(doc):
    """Generates the second page of Table of Contents matching page vii of reference PDF."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_borders(table)

    headers = ["CHAP.NO.", "TITLE", "PAGE.NO"]
    col_widths = [Inches(1.2), Inches(3.8), Inches(1.0)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(title)
        r.font.name = FONT_NAME
        r.font.size = Pt(10.5)
        r.font.bold = True

    entries = [
        ("4", "METHODOLOGY", "18"),
        ("", "4.1   EXISTING METHODOLOGY", "18"),
        ("", "4.2   PROPOSED METHODOLOGY", "19"),
        ("", "4.3   ADVANTAGES", "20"),
        ("", "4.4   DISADVANTAGES", "21"),
        ("", "4.5   APPLICATIONS", "21"),
        ("5", "SOFTWARE IMPLEMENTATION", "22"),
        ("", "5.1   OVERVIEW OF THE ARDUINO PLATFORM", "22"),
        ("", "5.2   PYTHON EDGE GATEWAY AND ML PIPELINE", "22"),
        ("", "5.3   IOT CONNECTIVITY AND TELEGRAM CLOUD MESSAGING", "22"),
        ("", "5.4   CONFIGURATION OF ARDUINO IDE AND SENSORS", "22"),
        ("6", "WORKING AND PROGRAM", "23"),
        ("", "6.1   WORKING", "23"),
        ("", "6.2   FLOWCHART", "24"),
        ("", "6.3   PROGRAM", "25"),
        ("", "6.4   TEST RESULTS", "29"),
        ("7", "CONCLUSION AND FUTURE WORK", "32"),
        ("", "7.1   CONCLUSION", "32"),
        ("", "7.2   FUTURE SCOPE", "32"),
        ("", "COST ESTIMATION", "33"),
        ("", "SUSTAINABILITY DEVELOPMENT GOALS (SDG)", "34"),
        ("", "REFERENCES", "36"),
    ]

    for chap, title, page in entries:
        row = table.add_row()
        cells = row.cells
        for i in range(3):
            cells[i].width = col_widths[i]

        p0 = cells[0].paragraphs[0]
        format_cell_p(p0, WD_ALIGN_PARAGRAPH.CENTER)
        r0 = p0.add_run(chap)
        r0.font.name = FONT_NAME
        r0.font.size = Pt(9.5)
        r0.font.bold = bool(chap)

        p1 = cells[1].paragraphs[0]
        format_cell_p(p1, WD_ALIGN_PARAGRAPH.LEFT)
        r1 = p1.add_run(title)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(9.5)
        r1.font.bold = bool(chap) or title in ["COST ESTIMATION", "SUSTAINABILITY DEVELOPMENT GOALS (SDG)", "REFERENCES"]

        p2 = cells[2].paragraphs[0]
        format_cell_p(p2, WD_ALIGN_PARAGRAPH.CENTER)
        r2 = p2.add_run(page)
        r2.font.name = FONT_NAME
        r2.font.size = Pt(9.5)
        r2.font.bold = bool(chap)

def build_list_of_figures_table(doc):
    """Builds the List of Figures table matching page viii of reference PDF."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_borders(table)

    headers = ["FIGURE NO", "NAME OF FIGURE", "PAGE NO"]
    col_widths = [Inches(1.4), Inches(3.6), Inches(1.0)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(title)
        r.font.name = FONT_NAME
        r.font.size = Pt(10.5)
        r.font.bold = True

    figures = [
        ("3.1", "Block Diagram of Proposed Method", "11"),
        ("3.2", "Circuit Diagram of Proposed Method", "13"),
        ("3.3", "Arduino Nano Microcontroller", "14"),
        ("3.4", "Solar Photovoltaic Panel", "15"),
        ("3.5", "18650 Lithium-ion Battery Cell", "16"),
        ("3.6", "ACS712 Current Sensor Module", "16"),
        ("3.7", "LM35 Precision Temperature Sensor", "17"),
        ("3.8", "IoT Edge Gateway and Cloud Communication", "17"),
        ("6.1", "Hardware Kit", "23"),
        ("6.2", "Flowchart", "24"),
        ("6.3", "Photo of Working Model", "30"),
    ]


    for f_no, f_title, page in figures:
        row = table.add_row()
        cells = row.cells
        for i in range(3):
            cells[i].width = col_widths[i]

        p0 = cells[0].paragraphs[0]
        format_cell_p(p0, WD_ALIGN_PARAGRAPH.CENTER)
        r0 = p0.add_run(f_no)
        r0.font.name = FONT_NAME
        r0.font.size = Pt(9.5)

        p1 = cells[1].paragraphs[0]
        format_cell_p(p1, WD_ALIGN_PARAGRAPH.LEFT)
        r1 = p1.add_run(f_title)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(9.5)

        p2 = cells[2].paragraphs[0]
        format_cell_p(p2, WD_ALIGN_PARAGRAPH.CENTER)
        r2 = p2.add_run(page)
        r2.font.name = FONT_NAME
        r2.font.size = Pt(9.5)

def build_test_table_1(doc):
    """Table 1: Battery Voltage & Load Current Monitoring (matches Table 1 on page 35 of sample PDF)."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("Table 1: Battery Voltage and Load Current Monitoring")
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(10.5)
    r_title.font.bold = True

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_borders(table)

    headers = ["S.No.", "Test Condition", "Display Status", "Operating Mode"]
    widths = [Inches(0.8), Inches(2.0), Inches(1.8), Inches(1.4)]

    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.width = widths[i]
        p = c.paragraphs[0]
        format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        r.font.name = FONT_NAME
        r.font.size = Pt(9.5)
        r.font.bold = True

    data = [
        ("1", "Load Disconnected (0 mA)", "Battery Idle (3.88 V)", "Idle Standby"),
        ("2", "DC Motor Active (480 mA)", "Discharging (3.82 V)", "Active Discharging"),
        ("3", "Solar Charging (-450 mA)", "Solar Inflow (3.95 V)", "Active Charging"),
    ]

    for row_data in data:
        row = table.add_row()
        for i, val in enumerate(row_data):
            c = row.cells[i]
            c.width = widths[i]
            p = c.paragraphs[0]
            format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(9)

def build_test_table_2(doc):
    """Table 2: Solar Panel Generation and Soiling Detection (matches Table 2 on page 35 of sample PDF)."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("Table 2: Solar Generation and Optical Soiling Diagnosis")
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(10.5)
    r_title.font.bold = True

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_borders(table)

    headers = ["S.No.", "Solar Exposure State", "Cleanliness", "Diagnostic Output"]
    widths = [Inches(0.8), Inches(1.8), Inches(1.4), Inches(2.0)]

    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.width = widths[i]
        p = c.paragraphs[0]
        format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        r.font.name = FONT_NAME
        r.font.size = Pt(9.5)
        r.font.bold = True

    data = [
        ("1", "Direct Sunlight (Clear Glass)", "Clean (98%)", "Optimal Harvest (Full Profile)"),
        ("2", "Dust Accumulation / Soiling", "Soiled (42%)", "Dust Alert Sent (-58% Loss)"),
        ("3", "Dusty Glass with Rain Forecast", "Soiled (42%)", "Rain Washing Expected (Wait)"),
    ]

    for row_data in data:
        row = table.add_row()
        for i, val in enumerate(row_data):
            c = row.cells[i]
            c.width = widths[i]
            p = c.paragraphs[0]
            format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(9)

def build_test_table_3(doc):
    """Table 3: IoT Cloud Gateway, ML Prognostics & Telegram Messaging (matches Table 3 on page 36 of PDF)."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(3)
    r_title = p_title.add_run("Table 3: Edge Machine Learning Prognostics and Telegram Messaging")
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(11)
    r_title.font.bold = True

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_borders(table)

    headers = ["S.No.", "Operating Condition", "Gateway ML & Hybrid SoC", "Telegram Result"]
    widths = [Inches(0.8), Inches(1.8), Inches(1.8), Inches(1.6)]

    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.width = widths[i]
        p = c.paragraphs[0]
        format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        r.font.name = FONT_NAME
        r.font.size = Pt(10)
        r.font.bold = True

    data = [
        ("1", "Normal Discharging with Sunlight", "SoH: 94%, Backup: 6h 40m", "Hourly 7-bullet report sent"),
        ("2", "Overcast Sky / Rain Forecast", "SoH: 92%, Rain Prob: 85%", "Power conservation advice"),
        ("3", "User sends '/report' command", "Instant telemetry query", "Layman diagnostic summary"),
        ("4", "User asks 'Can I run 10W fan?'", "Calculates Wh vs load", "Exact hours calculation returned"),
    ]

    for row_data in data:
        row = table.add_row()
        for i, val in enumerate(row_data):
            c = row.cells[i]
            c.width = widths[i]
            p = c.paragraphs[0]
            format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(9.5)

def build_cost_estimation_table(doc):
    """Builds the Cost Estimation table matching page 38 of reference PDF."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_borders(table)

    headers = ["S.NO.", "COMPONENTS", "PRICE(IN RUPEES)"]
    widths = [Inches(1.2), Inches(3.6), Inches(1.4)]

    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.width = widths[i]
        p = c.paragraphs[0]
        format_cell_p(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.font.bold = True

    components = [
        ("1", "Arduino Nano V3.0 (ATmega328P Board)", "280"),
        ("2", "Rock Light 3W / 6V Polycrystalline Solar Panel", "320"),
        ("3", "18650 3.7V 2200mAh Lithium-ion Cell & Holder", "190"),
        ("4", "ACS712-05A Hall-Effect Current Sensor Module", "140"),
        ("5", "LM35 Precision Analog Temperature Sensor IC", "60"),
        ("6", "Precision Metal Film Resistors (Voltage Dividers)", "20"),
        ("7", "Breadboard and Jumper Wires Kit", "90"),
        ("8", "Mini-USB Communication / Power Cable", "60"),
        ("9", "DC Demonstration Motor Load", "70"),
        ("10", "Edge Computing & Cloud Telegram Gateway Interface", "250"),
    ]

    for s_no, comp, price in components:
        row = table.add_row()
        cells = row.cells
        for i in range(3):
            cells[i].width = widths[i]

        p0 = cells[0].paragraphs[0]
        format_cell_p(p0, WD_ALIGN_PARAGRAPH.CENTER)
        r0 = p0.add_run(s_no)
        r0.font.name = FONT_NAME
        r0.font.size = Pt(10)

        p1 = cells[1].paragraphs[0]
        format_cell_p(p1, WD_ALIGN_PARAGRAPH.LEFT)
        r1 = p1.add_run(comp)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(10)

        p2 = cells[2].paragraphs[0]
        format_cell_p(p2, WD_ALIGN_PARAGRAPH.CENTER)
        r2 = p2.add_run(price)
        r2.font.name = FONT_NAME
        r2.font.size = Pt(10)

    # Total Row
    total_row = table.add_row()
    t_cells = total_row.cells
    for i in range(3):
        t_cells[i].width = widths[i]

    p_lbl = t_cells[1].paragraphs[0]
    format_cell_p(p_lbl, WD_ALIGN_PARAGRAPH.RIGHT)
    r_lbl = p_lbl.add_run("TOTAL")
    r_lbl.font.name = FONT_NAME
    r_lbl.font.size = Pt(11)
    r_lbl.font.bold = True

    p_val = t_cells[2].paragraphs[0]
    format_cell_p(p_val, WD_ALIGN_PARAGRAPH.CENTER)
    r_val = p_val.add_run("1480/-")
    r_val.font.name = FONT_NAME
    r_val.font.size = Pt(11)
    r_val.font.bold = True
