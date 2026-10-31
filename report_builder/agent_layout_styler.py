"""
agent_layout_styler.py
======================
Configures Word document geometry, styles, headers, footers, and XML field codes
to precisely replicate the typographic hierarchy and layout of the reference PDF.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
BODY_FONT_SIZE = Pt(12)
CHAPTER_FONT_SIZE = Pt(14)
HEADING_FONT_SIZE = Pt(12)
CAPTION_FONT_SIZE = Pt(11)

def configure_document_styles(doc):
    """Sets standard margins, font defaults, and line spacing on the document."""
    # Set default style
    normal_style = doc.styles['Normal']
    normal_style.font.name = FONT_NAME
    normal_style.font.size = BODY_FONT_SIZE
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Configure Section 1 (Preliminary Pages: Abstract to List of Figures)
    sec1 = doc.sections[0]
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.25)   # Binding margin
    sec1.right_margin = Inches(1.0)

    # Set page numbering to lowercase Roman, starting at v (5)
    sectPr1 = sec1._sectPr
    pgNumType1 = OxmlElement('w:pgNumType')
    pgNumType1.set(qn('w:fmt'), 'lowerRoman')
    pgNumType1.set(qn('w:start'), '5')
    sectPr1.append(pgNumType1)

    # Setup footer for Section 1
    footer1 = sec1.footer
    p1 = footer1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run()
    r1.font.name = FONT_NAME
    r1.font.size = Pt(11)
    _add_page_number_field(r1)

def add_body_section(doc):
    """Creates Section 2 for Main Chapters with Arabic (1, 2, 3...) page numbering."""
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.top_margin = Inches(1.0)
    sec2.bottom_margin = Inches(1.0)
    sec2.left_margin = Inches(1.25)
    sec2.right_margin = Inches(1.0)

    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False

    # Start Arabic page numbers at 1 (matches List of Abbreviations = 1, Chapter 1 = 2)
    sectPr2 = sec2._sectPr
    pgNumType2 = OxmlElement('w:pgNumType')
    pgNumType2.set(qn('w:fmt'), 'decimal')
    pgNumType2.set(qn('w:start'), '1')
    sectPr2.append(pgNumType2)

    footer2 = sec2.footer
    p2 = footer2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run()
    r2.font.name = FONT_NAME
    r2.font.size = Pt(11)
    _add_page_number_field(r2)
    return sec2

def _add_page_number_field(run):
    """Appends dynamic Word XML PAGE field."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def add_chapter_header(doc, chapter_num, chapter_title):
    """Adds centered, bold, uppercase chapter title matching the PDF."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    
    r1 = p.add_run(f"CHAPTER {chapter_num}\n")
    r1.font.name = FONT_NAME
    r1.font.size = CHAPTER_FONT_SIZE
    r1.font.bold = True
    
    r2 = p.add_run(chapter_title.upper())
    r2.font.name = FONT_NAME
    r2.font.size = CHAPTER_FONT_SIZE
    r2.font.bold = True
    return p

def add_unnum_title(doc, title_text):
    """Adds standalone centered bold title (e.g. ABSTRACT, REFERENCES, COST ESTIMATION)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.3
    
    r = p.add_run(title_text.upper())
    r.font.name = FONT_NAME
    r.font.size = CHAPTER_FONT_SIZE
    r.font.bold = True
    return p

def add_heading_1(doc, sec_num, sec_title):
    """Adds 1.1 GENERAL style heading (Bold, uppercase, left-aligned)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    
    r = p.add_run(f"{sec_num}  {sec_title.upper()}")
    r.font.name = FONT_NAME
    r.font.size = HEADING_FONT_SIZE
    r.font.bold = True
    return p

def add_heading_2(doc, subsec_num, subsec_title):
    """Adds 3.2.1 Arduino Nano style subsection heading."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    
    r = p.add_run(f"{subsec_num}  {subsec_title}")
    r.font.name = FONT_NAME
    r.font.size = HEADING_FONT_SIZE
    r.font.bold = True
    return p

def add_paragraph(doc, text):
    """Adds a standard justified 1.5 spaced paragraph."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = BODY_FONT_SIZE
    return p

def add_bullet_point(doc, text, bold_prefix=None):
    """Adds a bullet point with exact 1.5 line spacing and Times New Roman font."""
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = BODY_FONT_SIZE
        r_pre.font.bold = True
        
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = BODY_FONT_SIZE
    return p

def add_caption(doc, caption_text):
    """Adds a centered, bold figure or table caption matching the PDF."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    
    r = p.add_run(caption_text)
    r.font.name = FONT_NAME
    r.font.size = CAPTION_FONT_SIZE
    r.font.bold = True
    return p
